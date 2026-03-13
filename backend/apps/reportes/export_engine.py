"""
Data export engine — CSV and XLSX generation with Rust acceleration.

Dispatches to gravitea_rust for performance, falls back to
Python csv stdlib + openpyxl when the Rust extension is unavailable.
"""
import csv
import io
import logging

logger = logging.getLogger(__name__)

try:
    from gravitea_rust import (
        generate_csv as _rust_generate_csv,
        generate_xlsx as _rust_generate_xlsx,
    )
    _USE_RUST_EXPORT = True
except (ImportError, OSError):
    _USE_RUST_EXPORT = False
    logger.warning(
        "gravitea_rust export not available — using Python fallback"
    )


def generate_csv(
    rows: list[dict[str, str]],
    headers: list[str],
) -> bytes:
    """Generate CSV bytes with UTF-8 BOM and ordered columns.

    Args:
        rows: List of dicts mapping column names to string values.
        headers: Ordered list of column names for output order.

    Returns:
        CSV file content as bytes, prefixed with UTF-8 BOM.
    """
    if _USE_RUST_EXPORT:
        return bytes(_rust_generate_csv(rows, headers))

    # Python fallback
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row.get(h, "") for h in headers)
    csv_str = output.getvalue()
    return b"\xEF\xBB\xBF" + csv_str.encode("utf-8")


def generate_xlsx(
    rows: list[dict[str, str]],
    headers: list[str],
    column_widths: dict[str, float] | None = None,
) -> bytes:
    """Generate XLSX bytes with numeric detection and optional column widths.

    Args:
        rows: List of dicts mapping column names to string values.
        headers: Ordered list of column names for output order.
        column_widths: Optional dict of column name -> width in character units.

    Returns:
        XLSX file content as bytes.
    """
    if _USE_RUST_EXPORT:
        return bytes(_rust_generate_xlsx(rows, headers, column_widths))

    # Python fallback using openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    bold = Font(bold=True)

    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold

    # Write data with numeric detection
    for row_idx, row in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row.get(header, "")
            try:
                ws.cell(row=row_idx, column=col_idx, value=float(value))
            except (ValueError, TypeError):
                ws.cell(row=row_idx, column=col_idx, value=value)

    # Apply column widths
    if column_widths:
        for col_idx, header in enumerate(headers, 1):
            if header in column_widths:
                col_letter = ws.cell(row=1, column=col_idx).column_letter
                ws.column_dimensions[col_letter].width = column_widths[header]

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
